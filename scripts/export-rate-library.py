"""Export the rate library from the GRB pricing-tool workbook to JSON.

Usage:
    python3 scripts/export-rate-library.py "GRB Home Improvements - Pricing Tool.xlsx"

Reads the 'Rates - Labour & Materials' and 'Rates - Benchmarks & Fees' tabs
and writes src/data/rate-library.json. Re-run whenever the Excel master changes;
never hand-edit the output.
"""

import json
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "src" / "data" / "rate-library.json"


def rows(ws, start, stop=None):
    for r in ws.iter_rows(min_row=start, max_row=stop, values_only=True):
        if r[0] in (None, ""):
            if stop is None:
                return
            continue
        yield r


def section(ws, title_prefix):
    """Row index of the first data row after a section header + column header row."""
    for r in ws.iter_rows(min_col=1, max_col=1):
        cell = r[0]
        if isinstance(cell.value, str) and cell.value.startswith(title_prefix):
            return cell.row + 2
    raise KeyError(f"section not found: {title_prefix}")


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    lm = wb["Rates - Labour & Materials"]
    header = [c.value for c in lm[3]]
    local_col = next((i for i, h in enumerate(header) if isinstance(h, str) and h.strip().lower().startswith("local")), None)
    items = []
    for r in rows(lm, 4):
        cat, item, unit, rate, low, high, src = r[:7]
        entry = {"category": cat, "item": item, "unit": unit, "rate": rate, "low": low, "high": high, "source": src}
        # Optional "Local actual" column (Yes/No): local prices skip the region/inflation uplift.
        if local_col is not None and str(r[local_col] or "").strip().lower() in ("yes", "y", "true", "1"):
            entry["local"] = True
        items.append(entry)

    bf = wb["Rates - Benchmarks & Fees"]
    benchmarks = {
        r[0]: {"Budget": [r[1], r[2]], "Standard": [r[3], r[4]], "Premium": [r[5], r[6]], "notes": r[7]}
        for r in rows(bf, 4)
    }
    fees = [
        {"item": r[0], "low": r[1], "high": r[2], "unit": r[3], "source": r[4]}
        for r in rows(bf, section(bf, "FEES & PRELIMINARIES"))
    ]
    roofs = {r[0]: [r[1], r[2]] for r in rows(bf, section(bf, "ROOF TYPE MODIFIER"))}
    coverings = {
        r[0]: {"low": r[1], "high": r[2], "notes": r[4]}
        for r in rows(bf, section(bf, "ROOF COVERINGS"))
    }
    plant = [
        {"item": r[0], "day": r[1], "week": r[2], "notes": r[4]}
        for r in rows(bf, section(bf, "PLANT & TOOL HIRE"))
    ]

    data = {
        "source": Path(path).name,
        "items": items,
        "benchmarks": benchmarks,
        "fees": fees,
        "roofs": roofs,
        "roofCoverings": coverings,
        "plantHire": plant,
    }
    OUT.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    counts = {k: len(v) for k, v in data.items() if k != "source"}
    print(f"wrote {OUT.relative_to(Path.cwd()) if OUT.is_relative_to(Path.cwd()) else OUT}: {counts}")


if __name__ == "__main__":
    main(sys.argv[1])
